import os
import re
from docx import Document
from docx.shared import Inches

def natural_sort_key(s):
    """
    Splits strings to sort numbers properly.
    Ensures 'image2.png' comes before 'image10.png'.
    """
    return [int(text) if text.isdigit() else text.lower() for text in re.split('([0-9]+)', s)]

def compile_images_to_word():
    # 1. Automatically detect the folder this script is saved in
    base_directory = os.path.dirname(os.path.abspath(__file__))
    print(f"Scanning folder: {base_directory} ...\n")

    # 2. Setup the Word document
    doc = Document()
    doc.add_heading('Compiled Project Images', 0)

    # 3. Get the subfolders inside this directory
    folders = [f for f in os.listdir(base_directory) if os.path.isdir(os.path.join(base_directory, f))]
    folders.sort()

    images_processed = False

    for folder in folders:
        folder_path = os.path.join(base_directory, folder)
        
        # Identify valid image files
        valid_extensions = {".png", ".jpg", ".jpeg", ".gif", ".bmp"}
        images = [f for f in os.listdir(folder_path) if os.path.splitext(f)[1].lower() in valid_extensions]
        
        if not images:
            continue # Skip folders that don't have images
            
        images_processed = True
        doc.add_heading(folder, level=1)
        images.sort(key=natural_sort_key)

        # Insert images
        for img in images:
            img_path = os.path.join(folder_path, img)
            try:
                doc.add_picture(img_path, width=Inches(6.0))
            except Exception as e:
                print(f"Warning: Could not add {img_path}. Error: {e}")

    if not images_processed:
        print("No images found in any subfolders. Make sure your image folders are extracted here.")
        return

    # 4. Save the file right next to the script
    output_filename = os.path.join(base_directory, "Lawyer_App_Screenshots.docx")
    doc.save(output_filename)
    
    print(f"SUCCESS! Your Word document has been generated.")
    print(f"You can find it here: {output_filename}")

if __name__ == "__main__":
    compile_images_to_word()
