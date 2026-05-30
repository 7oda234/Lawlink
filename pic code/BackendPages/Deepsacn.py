import os
import re
from docx import Document
from docx.shared import Inches

def natural_sort_key(s):
    """
    Splits strings to sort numbers properly.
    Ensures 'image2.png' comes before 'image10.png'.
    """
    return [int(text) if text.isdigit() else text.lower() for text in re.split('([0-9]+)', s)]

def compile_images_to_word():
    # 1. Automatically detect the folder this script is saved in
    base_directory = os.path.dirname(os.path.abspath(__file__))
    print(f"Deep scanning folder and all subfolders in: {base_directory} ...\n")

    # 2. Setup the Word document
    doc = Document()
    doc.add_heading('Compiled Project Images', 0)

    images_processed = False

    # 3. os.walk() automatically digs into every single subfolder
    for root_folder, _, files in os.walk(base_directory):
        
        # Identify valid image files in the current folder
        valid_extensions = {".png", ".jpg", ".jpeg", ".gif", ".bmp"}
        images = [f for f in files if os.path.splitext(f)[1].lower() in valid_extensions]
        
        if not images:
            continue # Skip folders that don't have images
            
        images_processed = True
        
        # Get the name of the deepest folder we are currently in to use as the heading
        folder_name = os.path.basename(root_folder)
        doc.add_heading(folder_name, level=1)
        
        # Sort images naturally
        images.sort(key=natural_sort_key)

        # Insert images
        for img in images:
            img_path = os.path.join(root_folder, img)
            try:
                doc.add_picture(img_path, width=Inches(6.0))
            except Exception as e:
                print(f"Warning: Could not add {img_path}. Error: {e}")

    if not images_processed:
        print("No images found in any subfolders. Make sure your image folders are extracted here.")
        return

    # 4. Save the file right next to the script
    output_filename = os.path.join(base_directory, "Lawyer_App_Screenshots.docx")
    doc.save(output_filename)
    
    print(f"SUCCESS! Your Word document has been generated.")
    print(f"You can find it here: {output_filename}")

if __name__ == "__main__":
    compile_images_to_word()
